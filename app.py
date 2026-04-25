from __future__ import annotations

import os
import re
import secrets
import logging
import csv
import io
import shutil
import sqlite3
import subprocess
import sys
from datetime import datetime
from functools import wraps
from pathlib import Path
from typing import Any

from docx import Document
from docx.oxml.ns import qn
from flask import Flask, Response, g, jsonify, redirect, render_template, request, send_from_directory, session, url_for
from werkzeug.security import check_password_hash, generate_password_hash

try:
    from docx2pdf import convert as docx2pdf_convert
except ImportError:  # pragma: no cover - optional dependency at runtime
    docx2pdf_convert = None


BASE_DIR = Path(__file__).resolve().parent
TEMPLATE_PATH = BASE_DIR / "template.docx"
OUTPUT_DIR = BASE_DIR / "output"
OUTPUT_DOCX = OUTPUT_DIR / "report.docx"
OUTPUT_PDF = OUTPUT_DIR / "report.pdf"
DATABASE_PATH = BASE_DIR / "mallika_auth.db"
DOCX2PDF_TIMEOUT_SECONDS = 60
PHONE_PATTERN = re.compile(r"^\d{10}$")
LIBREOFFICE_CANDIDATES = (
    r"C:\Program Files\LibreOffice\program\soffice.exe",
    r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
)


app = Flask(__name__)
app.config["JSON_SORT_KEYS"] = False
app.config["SECRET_KEY"] = os.getenv("APP_SECRET_KEY", "mallika-change-this-secret-key")
app.config["REGISTRATION_SECRET"] = os.getenv("REGISTRATION_SECRET", "")


def get_db() -> sqlite3.Connection:
    if "db" not in g:
        database = sqlite3.connect(DATABASE_PATH)
        database.row_factory = sqlite3.Row
        g.db = database
    return g.db


@app.teardown_appcontext
def close_db(_exc: BaseException | None) -> None:
    database = g.pop("db", None)
    if database is not None:
        database.close()


def init_db() -> None:
    db = get_db()
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS users (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT NOT NULL UNIQUE,
            password_hash TEXT NOT NULL,
            created_at TEXT NOT NULL,
            is_admin INTEGER NOT NULL DEFAULT 0
        )
        """
    )
    db.execute(
        """
        CREATE TABLE IF NOT EXISTS report_submissions (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            user_id INTEGER NOT NULL,
            username TEXT NOT NULL,
            created_at TEXT NOT NULL,
            patient_name TEXT NOT NULL,
            patient_number TEXT,
            patient_address TEXT,
            patient_age TEXT,
            patient_gender TEXT,
            height_cm TEXT,
            weight_kg TEXT,
            bmi_value TEXT,
            pulse TEXT,
            bp TEXT,
            rr TEXT,
            temp TEXT,
            blood_sugar TEXT,
            impression TEXT,
            report_docx TEXT NOT NULL,
            report_pdf TEXT NOT NULL,
            FOREIGN KEY(user_id) REFERENCES users(id)
        )
        """
    )
    user_columns = {row["name"] for row in db.execute("PRAGMA table_info(users)").fetchall()}
    if "is_admin" not in user_columns:
        db.execute("ALTER TABLE users ADD COLUMN is_admin INTEGER NOT NULL DEFAULT 0")

    report_columns = {row["name"] for row in db.execute("PRAGMA table_info(report_submissions)").fetchall()}
    if "username" not in report_columns:
        db.execute("ALTER TABLE report_submissions ADD COLUMN username TEXT NOT NULL DEFAULT ''")

    if not db.execute("SELECT 1 FROM users WHERE is_admin = 1 LIMIT 1").fetchone():
        first_user = db.execute("SELECT id FROM users ORDER BY id LIMIT 1").fetchone()
        if first_user:
            db.execute("UPDATE users SET is_admin = 1 WHERE id = ?", (first_user["id"],))
    db.commit()


def get_or_create_csrf_token() -> str:
    token = session.get("_csrf_token")
    if not token:
        token = secrets.token_hex(32)
        session["_csrf_token"] = token
    return token


def is_valid_csrf(token: str | None) -> bool:
    session_token = session.get("_csrf_token")
    return bool(token and session_token and secrets.compare_digest(token, session_token))


def fetch_current_user() -> sqlite3.Row | None:
    user_id = session.get("user_id")
    if not user_id:
        return None
    return get_db().execute(
        "SELECT id, username, created_at, is_admin FROM users WHERE id = ?",
        (user_id,),
    ).fetchone()


@app.before_request
def load_current_user() -> None:
    g.current_user = fetch_current_user()


@app.context_processor
def inject_template_helpers() -> dict[str, Any]:
    return {
        "csrf_token": get_or_create_csrf_token,
        "current_user": getattr(g, "current_user", None),
        "registration_enabled": bool(app.config["REGISTRATION_SECRET"]),
    }


def login_required(view):
    @wraps(view)
    def wrapped_view(*args, **kwargs):
        if getattr(g, "current_user", None) is None:
            if request.is_json or request.path == "/generate-report":
                return jsonify({"status": "error", "message": "Authentication required"}), 401
            return redirect(url_for("login"))
        return view(*args, **kwargs)

    return wrapped_view


def admin_required(view):
    @wraps(view)
    @login_required
    def wrapped_view(*args, **kwargs):
        if not getattr(g, "current_user", None) or not g.current_user["is_admin"]:
            return "Admin access required", 403
        return view(*args, **kwargs)

    return wrapped_view


def stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return format(value, "g")
    return str(value).strip()


def safe_float(value: Any) -> float | None:
    text = stringify(value)
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def calculate_bmi_value(height_cm: Any, weight_kg: Any) -> str:
    height = safe_float(height_cm)
    weight = safe_float(weight_kg)
    if height is None or weight is None or height <= 0 or weight <= 0:
        return ""

    bmi = weight / ((height / 100) ** 2)
    return f"{bmi:.2f}"


def normalize_payload(payload: dict[str, Any]) -> dict[str, str]:
    number = stringify(payload.get("number"))
    if number and not PHONE_PATTERN.fullmatch(number):
        raise ValueError("number must contain exactly 10 digits")

    return {
        "name": stringify(payload.get("name")),
        "age": stringify(payload.get("age")),
        "gender": stringify(payload.get("gender")),
        "number": number,
        "address": stringify(payload.get("address")),
        "height_cm": stringify(payload.get("height_cm")),
        "weight_kg": stringify(payload.get("weight_kg")),
        "pulse": stringify(payload.get("pulse")),
        "bp": stringify(payload.get("bp")),
        "rr": stringify(payload.get("rr")),
        "temp": stringify(payload.get("temp")),
        "blood_sugar": stringify(payload.get("blood_sugar")),
        "impression": stringify(payload.get("impression")),
    }


def build_placeholder_map(payload: dict[str, str]) -> dict[str, str]:
    return {
        "{{date}}": datetime.now().strftime("%d-%m-%Y"),
        "{{name}}": payload["name"],
        "{{age}}": payload["age"],
        "{{gender}}": payload["gender"],
        "{{number}}": payload["number"],
        "{{add}}": payload["address"],
        "{{ht}}": payload["height_cm"],
        "{{wt}}": payload["weight_kg"],
        "{{cat}}": calculate_bmi_value(payload["height_cm"], payload["weight_kg"]),
        "{{pul}}": payload["pulse"],
        "{{bp}}": payload["bp"],
        "{{rr}}": payload["rr"],
        "{{temp}}": payload["temp"],
        "{{bs}}": payload["blood_sugar"],
        "{{imp}}": payload["impression"],
    }


def replace_placeholders_in_text_nodes(text_nodes: list[Any], replacements: dict[str, str]) -> None:
    if not text_nodes:
        return

    for placeholder, replacement in replacements.items():
        while True:
            full_text = "".join(node.text or "" for node in text_nodes)
            start_index = full_text.find(placeholder)
            if start_index < 0:
                break

            end_index = start_index + len(placeholder)
            node_matches: list[tuple[int, int, int]] = []
            cursor = 0

            for node_index, node in enumerate(text_nodes):
                node_text = node.text or ""
                node_start = cursor
                node_end = cursor + len(node_text)
                cursor = node_end

                if node_end <= start_index or node_start >= end_index:
                    continue

                node_matches.append(
                    (
                        node_index,
                        max(start_index - node_start, 0),
                        min(end_index - node_start, len(node_text)),
                    )
                )

            if not node_matches:
                break

            first_node_index, first_start, _ = node_matches[0]
            last_node_index, _, last_end = node_matches[-1]

            if first_node_index == last_node_index:
                node = text_nodes[first_node_index]
                node_text = node.text or ""
                node.text = node_text[:first_start] + replacement + node_text[last_end:]
                continue

            first_node = text_nodes[first_node_index]
            last_node = text_nodes[last_node_index]
            first_text = first_node.text or ""
            last_text = last_node.text or ""

            first_node.text = first_text[:first_start] + replacement

            for node_index, _, _ in node_matches[1:-1]:
                text_nodes[node_index].text = ""

            last_node.text = last_text[last_end:]


def replace_placeholders_in_xml_root(root: Any, replacements: dict[str, str]) -> None:
    for paragraph in root.iter(qn("w:p")):
        text_nodes = [node for node in paragraph.iter(qn("w:t"))]
        replace_placeholders_in_text_nodes(text_nodes, replacements)


def replace_placeholders(document: Document, replacements: dict[str, str]) -> None:
    replace_placeholders_in_xml_root(document._element, replacements)

    for section in document.sections:
        replace_placeholders_in_xml_root(section.header._element, replacements)
        replace_placeholders_in_xml_root(section.footer._element, replacements)


def list_word_processes() -> list[str]:
    if os.name != "nt":
        return []

    result = subprocess.run(
        ["tasklist", "/FI", "IMAGENAME eq WINWORD.EXE", "/FO", "CSV", "/NH"],
        capture_output=True,
        check=False,
        creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        text=True,
    )
    if result.returncode != 0:
        return []

    lines = [line.strip() for line in result.stdout.splitlines() if line.strip()]
    return [line for line in lines if not line.startswith("INFO:")]


def find_libreoffice_binary() -> str | None:
    from_path = shutil.which("soffice") or shutil.which("libreoffice")
    if from_path:
        return from_path

    for candidate in LIBREOFFICE_CANDIDATES:
        if Path(candidate).exists():
            return candidate
    return None


def convert_with_word_com(docx_path: Path, pdf_path: Path) -> None:
    script = "\n".join(
        [
            "import sys",
            "import pythoncom",
            "from pathlib import Path",
            "from win32com.client import DispatchEx",
            "",
            "pythoncom.CoInitialize()",
            "word = DispatchEx('Word.Application')",
            "word.Visible = False",
            "word.DisplayAlerts = 0",
            "word.AutomationSecurity = 3",
            "doc = None",
            "src = str(Path(sys.argv[1]).resolve())",
            "dst = str(Path(sys.argv[2]).resolve())",
            "try:",
            "    doc = word.Documents.Open(",
            "        src,",
            "        ConfirmConversions=False,",
            "        ReadOnly=True,",
            "        AddToRecentFiles=False,",
            "        OpenAndRepair=True,",
            "    )",
            "    doc.ExportAsFixedFormat(",
            "        OutputFileName=dst,",
            "        ExportFormat=17,",
            "        OpenAfterExport=False,",
            "        OptimizeFor=0,",
            "        CreateBookmarks=0,",
            "    )",
            "finally:",
            "    if doc is not None:",
            "        doc.Close(False)",
            "    word.Quit()",
            "    pythoncom.CoUninitialize()",
        ]
    )
    command = [
        sys.executable,
        "-c",
        script,
        str(docx_path),
        str(pdf_path),
    ]
    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    try:
        result = subprocess.run(
            command,
            capture_output=True,
            check=False,
            creationflags=creationflags,
            text=True,
            timeout=DOCX2PDF_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(
            "Word COM export timed out. Close all Microsoft Word windows and try again."
        ) from exc

    if result.returncode != 0:
        details = result.stderr.strip() or result.stdout.strip() or "unknown Word COM error"
        raise RuntimeError(f"Word COM export failed: {details}")

    if not pdf_path.exists():
        raise RuntimeError("Word COM export finished without creating the PDF file")


def convert_with_docx2pdf(docx_path: Path, pdf_path: Path) -> None:
    if docx2pdf_convert is None:
        raise RuntimeError("docx2pdf is not installed")

    command = [
        sys.executable,
        "-c",
        "from docx2pdf import convert; import sys; convert(sys.argv[1], sys.argv[2])",
        str(docx_path),
        str(pdf_path),
    ]
    creationflags = getattr(subprocess, "CREATE_NO_WINDOW", 0)

    try:
        result = subprocess.run(
            command,
            capture_output=True,
            check=False,
            creationflags=creationflags,
            text=True,
            timeout=DOCX2PDF_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError(
            "docx2pdf timed out while automating Microsoft Word. "
            "Close any blocked Word windows or install LibreOffice for the fallback path."
        ) from exc

    if result.returncode != 0:
        details = result.stderr.strip() or result.stdout.strip() or "unknown docx2pdf error"
        raise RuntimeError(f"docx2pdf failed: {details}")

    if not pdf_path.exists():
        raise RuntimeError("docx2pdf finished without creating the PDF file")


def convert_docx_to_pdf(docx_path: Path, pdf_path: Path) -> None:
    conversion_errors: list[str] = []
    if pdf_path.exists():
        pdf_path.unlink()

    if os.name == "nt":
        try:
            convert_with_word_com(docx_path, pdf_path)
            return
        except Exception as exc:  # pragma: no cover - depends on local Office setup
            conversion_errors.append(str(exc))

        try:
            convert_with_docx2pdf(docx_path, pdf_path)
            return
        except Exception as exc:  # pragma: no cover - depends on local Office setup
            conversion_errors.append(str(exc))

    libreoffice_binary = find_libreoffice_binary()
    if libreoffice_binary:
        command = [
            libreoffice_binary,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(pdf_path.parent),
            str(docx_path),
        ]
        result = subprocess.run(command, capture_output=True, text=True, check=False)
        if result.returncode == 0 and pdf_path.exists():
            return

        stderr = result.stderr.strip() or result.stdout.strip() or "unknown LibreOffice error"
        conversion_errors.append(f"LibreOffice failed: {stderr}")
    else:
        conversion_errors.append("LibreOffice/soffice is not installed")

    if os.name == "nt":
        word_processes = list_word_processes()
        if word_processes:
            conversion_errors.append(
                f"Microsoft Word appears to still be open ({len(word_processes)} WINWORD process(es)). "
                "Close every Word window, then retry."
            )

    raise RuntimeError("PDF conversion failed. " + " | ".join(conversion_errors))


def build_output_paths() -> tuple[Path, Path]:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    return OUTPUT_DIR / f"report_{stamp}.docx", OUTPUT_DIR / f"report_{stamp}.pdf"


def refresh_latest_copies(source_docx: Path, source_pdf: Path) -> None:
    for source_path, target_path in ((source_docx, OUTPUT_DOCX), (source_pdf, OUTPUT_PDF)):
        try:
            shutil.copyfile(source_path, target_path)
        except OSError as exc:
            app.logger.warning("Could not refresh latest copy %s: %s", target_path.name, exc)


def generate_report_files(payload: dict[str, str]) -> tuple[Path, Path]:
    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(f"Template not found: {TEMPLATE_PATH}")

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_docx, output_pdf = build_output_paths()
    document = Document(str(TEMPLATE_PATH))
    replace_placeholders(document, build_placeholder_map(payload))
    document.save(str(output_docx))
    convert_docx_to_pdf(output_docx, output_pdf)
    refresh_latest_copies(output_docx, output_pdf)
    return output_docx, output_pdf


def record_report_submission(
    user: sqlite3.Row,
    payload: dict[str, str],
    output_docx: Path,
    output_pdf: Path,
) -> None:
    created_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
    db = get_db()
    db.execute(
        """
        INSERT INTO report_submissions (
            user_id,
            username,
            created_at,
            patient_name,
            patient_number,
            patient_address,
            patient_age,
            patient_gender,
            height_cm,
            weight_kg,
            bmi_value,
            pulse,
            bp,
            rr,
            temp,
            blood_sugar,
            impression,
            report_docx,
            report_pdf
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        (
            user["id"],
            user["username"],
            created_at,
            payload["name"],
            payload["number"],
            payload["address"],
            payload["age"],
            payload["gender"],
            payload["height_cm"],
            payload["weight_kg"],
            calculate_bmi_value(payload["height_cm"], payload["weight_kg"]),
            payload["pulse"],
            payload["bp"],
            payload["rr"],
            payload["temp"],
            payload["blood_sugar"],
            payload["impression"],
            output_docx.name,
            output_pdf.name,
        ),
    )
    db.commit()


def build_admin_summary() -> dict[str, Any]:
    db = get_db()
    total_users = db.execute("SELECT COUNT(*) AS count FROM users").fetchone()["count"]
    total_forms = db.execute("SELECT COUNT(*) AS count FROM report_submissions").fetchone()["count"]
    total_patients = db.execute(
        "SELECT COUNT(DISTINCT NULLIF(TRIM(patient_number), '')) AS count FROM report_submissions"
    ).fetchone()["count"]
    latest_rows = db.execute(
        """
        SELECT
            id,
            username,
            created_at,
            patient_name,
            patient_number,
            patient_address,
            patient_age,
            patient_gender,
            bmi_value
        FROM report_submissions
        ORDER BY id DESC
        LIMIT 100
        """
    ).fetchall()
    user_rows = db.execute(
        """
        SELECT
            username,
            is_admin,
            created_at,
            (
                SELECT COUNT(*)
                FROM report_submissions rs
                WHERE rs.user_id = users.id
            ) AS forms_count
        FROM users
        ORDER BY is_admin DESC, username COLLATE NOCASE ASC
        """
    ).fetchall()
    return {
        "total_users": total_users,
        "total_forms": total_forms,
        "total_patients": total_patients,
        "submissions": latest_rows,
        "users": user_rows,
    }


@app.route("/login", methods=["GET", "POST"])
def login():
    if getattr(g, "current_user", None) is not None:
        return redirect(url_for("index"))

    error = ""
    if request.method == "POST":
        if not is_valid_csrf(request.form.get("csrf_token")):
            error = "Session expired. Reload the page and try again."
        else:
            username = stringify(request.form.get("username"))
            password = request.form.get("password") or ""
            user = get_db().execute("SELECT * FROM users WHERE username = ?", (username,)).fetchone()

            if user and check_password_hash(user["password_hash"], password):
                session.clear()
                session["user_id"] = user["id"]
                session["username"] = user["username"]
                session["_csrf_token"] = secrets.token_hex(32)
                return redirect(url_for("index"))

            error = "Invalid username or password."

    return render_template("auth.html", mode="login", error=error)


@app.route("/register", methods=["GET", "POST"])
def register():
    if getattr(g, "current_user", None) is not None:
        return redirect(url_for("index"))

    if not app.config["REGISTRATION_SECRET"]:
        return (
            render_template(
                "auth.html",
                mode="register",
                error="Registration is disabled until REGISTRATION_SECRET is configured on the server.",
            ),
            503,
        )

    error = ""
    if request.method == "POST":
        if not is_valid_csrf(request.form.get("csrf_token")):
            error = "Session expired. Reload the page and try again."
        else:
            username = stringify(request.form.get("username"))
            password = request.form.get("password") or ""
            confirm_password = request.form.get("confirm_password") or ""
            registration_secret = request.form.get("registration_secret") or ""
            db = get_db()

            if registration_secret != app.config["REGISTRATION_SECRET"]:
                error = "Invalid registration key."
            elif not username:
                error = "Username is required."
            elif len(password) < 8:
                error = "Password must be at least 8 characters."
            elif password != confirm_password:
                error = "Passwords do not match."
            elif db.execute("SELECT 1 FROM users WHERE username = ?", (username,)).fetchone():
                error = "Username already exists."
            else:
                created_at = datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")
                cursor = db.execute(
                    "INSERT INTO users (username, password_hash, created_at, is_admin) VALUES (?, ?, ?, ?)",
                    (
                        username,
                        generate_password_hash(password),
                        created_at,
                        0 if db.execute("SELECT 1 FROM users LIMIT 1").fetchone() else 1,
                    ),
                )
                db.commit()
                session.clear()
                session["user_id"] = cursor.lastrowid
                session["username"] = username
                session["_csrf_token"] = secrets.token_hex(32)
                return redirect(url_for("index"))

    return render_template("auth.html", mode="register", error=error)


@app.post("/logout")
@login_required
def logout():
    if not is_valid_csrf(request.form.get("csrf_token")):
        return "Invalid security token", 400
    session.clear()
    return redirect(url_for("login"))


@app.post("/generate-report")
@login_required
def generate_report():
    if not is_valid_csrf(request.headers.get("X-CSRF-Token")):
        return jsonify({"status": "error", "message": "Invalid or missing security token"}), 403

    payload = request.get_json(silent=True)
    if payload is None or not isinstance(payload, dict):
        return (
            jsonify(
                {
                    "status": "error",
                    "message": "Request body must be valid JSON object",
                }
            ),
            400,
        )

    try:
        normalized_payload = normalize_payload(payload)
        docx_path, pdf_path = generate_report_files(normalized_payload)
        record_report_submission(g.current_user, normalized_payload, docx_path, pdf_path)
    except ValueError as exc:
        return jsonify({"status": "error", "message": str(exc)}), 400
    except FileNotFoundError as exc:
        app.logger.exception("Template error")
        return jsonify({"status": "error", "message": str(exc)}), 500
    except Exception as exc:  # pragma: no cover - integration/runtime path
        app.logger.exception("Failed to generate report")
        return jsonify({"status": "error", "message": str(exc)}), 500

    return jsonify({"status": "success", "pdf_url": f"/download/{pdf_path.name}"})


@app.get("/download/<path:filename>")
@login_required
def download_file(filename: str):
    safe_name = Path(filename).name
    file_path = OUTPUT_DIR / safe_name
    if not file_path.exists():
        return jsonify({"status": "error", "message": "File not found"}), 404
    return send_from_directory(OUTPUT_DIR, safe_name, as_attachment=True)


@app.get("/")
@login_required
def index():
    return render_template("index.html")


@app.get("/admin")
@admin_required
def admin_dashboard():
    return render_template("admin.html", **build_admin_summary())


@app.get("/admin/export.csv")
@admin_required
def admin_export_csv():
    rows = get_db().execute(
        """
        SELECT
            created_at,
            username,
            patient_name,
            patient_number,
            patient_address,
            patient_age,
            patient_gender,
            height_cm,
            weight_kg,
            bmi_value,
            pulse,
            bp,
            rr,
            temp,
            blood_sugar,
            impression,
            report_docx,
            report_pdf
        FROM report_submissions
        ORDER BY id DESC
        """
    ).fetchall()

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(
        [
            "generated_at",
            "submitted_by",
            "patient_name",
            "patient_number",
            "patient_address",
            "patient_age",
            "patient_gender",
            "height_cm",
            "weight_kg",
            "bmi_value",
            "pulse",
            "bp",
            "rr",
            "temp",
            "blood_sugar",
            "impression",
            "report_docx",
            "report_pdf",
        ]
    )
    for row in rows:
        writer.writerow([row[column] for column in row.keys()])

    filename = f"mallika-report-export-{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(
        buffer.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@app.get("/health")
def health():
    return jsonify(
        {
            "status": "ready",
            "message": "Mallika Hospital report service is ready",
        }
    )


with app.app_context():
    init_db()


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    host = os.getenv("FLASK_HOST", "127.0.0.1")
    port = int(os.getenv("PORT") or os.getenv("FLASK_PORT", "8000"))
    debug = os.getenv("FLASK_DEBUG", "1").strip().lower() in {"1", "true", "yes", "on"}
    app.run(host=host, port=port, debug=debug)
